#!/usr/bin/env python3
"""shapes.json -> quantity takeoff workbook (Summary / Takeoff / Assumptions / RFIs).

  build_xlsx.py shapes.json -o Takeoff.xlsx [--title "Project – Takeoff"] [--rfis rfis.json]

Takeoff tab: one row per shape with measured SF/LF/EA as blue inputs and CY as a live formula of the
depth cell, a yellow Unit Cost column and Extended $ formula so the GC's itemized form can be priced in place.
rfis.json (optional): [{"issue": "...", "ref": "C-410", "impact": "..."}]
"""
import argparse, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from common import load_shapes, quantities, primary

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF"); BLK = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
HDR = Font(name=F, size=10, bold=True, color="FFFFFF"); TITLE = Font(name=F, size=14, bold=True)
HFILL = PatternFill("solid", fgColor="1F3864"); SFILL = PatternFill("solid", fgColor="D9E1F2"); YFILL = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0.0;(#,##0.0);"-"'; CUR = '$#,##0;($#,##0);"-"'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("shapes"); ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--title", default="Quantity Takeoff"); ap.add_argument("--rfis")
    a = ap.parse_args()
    data = load_shapes(a.shapes); k = {s["id"]: s for s in data["sheets"]}
    wb = Workbook(); T = wb.active; T.title = "Takeoff"
    T["A1"] = a.title; T["A1"].font = TITLE
    T["A2"] = "Blue = measured from drawing (vector layers / traced); black = formula; yellow = enter unit cost. Depth/width cells drive CY/SF live."
    T["A2"].font = Font(name=F, size=9, italic=True)
    heads = ["ID", "Sheet", "Item", "Scope", "Type", "Count", "Measured", "Unit", "Width (ft)", "Depth (in)", "SF", "LF", "EA", "CY", "Unit Cost ($)", "Extended ($)", "Notes"]
    for j, h in enumerate(heads, 1):
        c = T.cell(4, j, h); c.font = HDR; c.fill = HFILL; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
    row = 5; scopes = {}
    order = list(data["scopes"].keys())
    shapes = sorted(data["shapes"], key=lambda s: (order.index(s["scope"]) if s["scope"] in order else 99, s["sheet"], s["tid"]))
    cur = None
    for s in shapes:
        if s["scope"] != cur:
            cur = s["scope"]; T.cell(row, 1, cur).font = BOLD
            for j in range(1, 18): T.cell(row, j).fill = SFILL; T.cell(row, j).border = BOX
            row += 1; scopes[cur] = [row, row]
        q = quantities(s, k[s["sheet"]]["ptPerFt"]); v, u = primary(s, q)
        base = v / (s.get("count") or 1)  # measured value for ONE instance; Count multiplies
        vals = [s["tid"], s["sheet"], s["name"], s["scope"], s["type"], s.get("count") or 1, round(base, 2), u, s.get("width") or 0, s.get("depth") or 0]
        for j, val in enumerate(vals, 1):
            c = T.cell(row, j, val); c.font = BLUE if j in (6, 7, 9, 10) else BLK; c.border = BOX
            if j in (6, 7, 9, 10): c.fill = YFILL if j in (9, 10) else PatternFill()
        r = row
        if s["type"] == "area":
            T.cell(r, 11, f"=F{r}*G{r}"); T.cell(r, 12, 0); T.cell(r, 13, 0); T.cell(r, 14, f"=K{r}*J{r}/12/27")
        elif s["type"] == "line":
            T.cell(r, 11, f"=L{r}*I{r}"); T.cell(r, 12, f"=F{r}*G{r}"); T.cell(r, 13, 0); T.cell(r, 14, f"=K{r}*J{r}/12/27")
        else:
            T.cell(r, 11, 0); T.cell(r, 12, 0); T.cell(r, 13, f"=F{r}*G{r}"); T.cell(r, 14, 0)
        uc = T.cell(r, 15); uc.fill = YFILL; uc.font = BLUE; uc.number_format = '$#,##0.00'
        T.cell(r, 16, f'=IF(AND(K{r}>0,N{r}=0),K{r},IF(N{r}>0,N{r},IF(L{r}>0,L{r},M{r})))*O{r}').number_format = CUR
        T.cell(r, 17, s.get("note", ""))
        for j in range(11, 18):
            T.cell(r, j).border = BOX; T.cell(r, j).font = BLK
            if j <= 14: T.cell(r, j).number_format = NUM
        T.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top"); T.cell(r, 17).alignment = Alignment(wrap_text=True, vertical="top")
        scopes[cur][1] = row; row += 1
    last = row - 1
    T.cell(row + 1, 3, "TOTAL").font = BOLD
    for col in "KLMNP":
        j = ord(col) - 64; c = T.cell(row + 1, j, f"=SUM({col}5:{col}{last})"); c.font = BOLD; c.number_format = CUR if col == "P" else NUM
    T.cell(row + 2, 3, "Extended $ uses CY if a depth is set, else SF, else LF, else EA – change the formula if you price a line differently.").font = Font(name=F, size=9, italic=True)
    for j, w in enumerate([10, 7, 48, 18, 7, 7, 10, 6, 9, 9, 10, 10, 7, 9, 12, 13, 50], 1):
        T.column_dimensions[get_column_letter(j)].width = w
    T.freeze_panes = "D5"; T.auto_filter.ref = f"A4:Q{last}"

    S = wb.create_sheet("Summary", 0); S["A1"] = a.title + " – scope summary"; S["A1"].font = TITLE
    for j, h in enumerate(["Scope", "Items", "SF", "LF", "EA", "CY"], 1):
        c = S.cell(3, j, h); c.font = HDR; c.fill = HFILL; c.border = BOX
    r = 4
    for sc, (r0, r1) in scopes.items():
        S.cell(r, 1, sc); S.cell(r, 2, r1 - r0 + 1)
        for j, col in zip(range(3, 7), "KLMN"):
            c = S.cell(r, j, f"=SUM(Takeoff!{col}{r0}:{col}{r1})"); c.number_format = NUM; c.font = Font(name=F, size=10, color="008000")
        for j in range(1, 7): S.cell(r, j).border = BOX
        r += 1
    S.cell(r, 1, "TOTAL").font = BOLD
    for j, col in zip(range(2, 7), "BCDEF"):
        c = S.cell(r, j, f"=SUM({col}4:{col}{r-1})"); c.font = BOLD; c.number_format = NUM
    S.cell(r + 1, 1, "Total priced (Takeoff tab)").font = BOLD; c = S.cell(r + 1, 6, f"=Takeoff!P{last+2}"); c.number_format = CUR; c.font = BOLD
    S.column_dimensions["A"].width = 26
    for col in "BCDEF": S.column_dimensions[col].width = 11

    A = wb.create_sheet("Assumptions"); A["A1"] = "Assumptions carried in the takeoff (fill in / confirm by RFI)"; A["A1"].font = TITLE
    rows = [("Sheet scale", ", ".join(f"{s['name']}: {s['ptPerFt']:.2f} pt/ft" for s in data["sheets"])),
            ("Area depths", "Depth (in) per line on the Takeoff tab – blank/0 means no CY computed"),
            ("Line widths", "Width (ft) per line on the Takeoff tab (trench, curb, patch) – drives SF/CY")]
    for i, (kk, vv) in enumerate(rows, 3):
        A.cell(i, 1, kk).font = BOLD; A.cell(i, 2, vv)
    A.column_dimensions["A"].width = 22; A.column_dimensions["B"].width = 110

    if a.rfis:
        Q = wb.create_sheet("RFIs"); Q["A1"] = "Discrepancies / items requiring clarification"; Q["A1"].font = TITLE
        for j, h in enumerate(["#", "Issue", "Sheet ref", "Takeoff impact", "Response"], 1):
            c = Q.cell(3, j, h); c.font = HDR; c.fill = HFILL; c.border = BOX
        for i, rf in enumerate(json.load(open(a.rfis)), 1):
            for j, val in enumerate([i, rf.get("issue", ""), rf.get("ref", ""), rf.get("impact", ""), ""], 1):
                c = Q.cell(i + 3, j, val); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
        Q.column_dimensions["B"].width = 80; Q.column_dimensions["C"].width = 18; Q.column_dimensions["D"].width = 40; Q.column_dimensions["E"].width = 30
    wb.calculation.fullCalcOnLoad = True
    wb.save(a.out); print(f"wrote {a.out}: {len(shapes)} items")


if __name__ == "__main__":
    main()
